import io, re
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from .models import ResumeProfile
from .taxonomy import SKILLS

SECTION_ALIASES = {
    "education": ["education", "academic background"],
    "experience": ["experience", "work experience", "professional experience", "employment"],
    "skills": ["skills", "technical skills", "core competencies"],
    "projects": ["projects", "selected projects"],
    "leadership": ["leadership", "activities", "extracurricular activities"],
    "coursework": ["coursework", "relevant coursework"],
    "certifications": ["certifications", "licenses"],
}

def extract_text(uploaded) -> tuple[str, list[str]]:
    name = getattr(uploaded, "name", "resume.txt").lower()
    data = uploaded.getvalue() if hasattr(uploaded, "getvalue") else uploaded.read()
    issues = []
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        pages = [p.extract_text() or "" for p in reader.pages]
        text = "\n".join(pages)
        if any(not p.strip() for p in pages): issues.append("One or more PDF pages produced little/no extractable text; scanned or graphical content may not parse.")
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        paras = [p.text for p in doc.paragraphs]
        table_text = [cell.text for t in doc.tables for row in t.rows for cell in row.cells]
        if doc.tables: issues.append("DOCX contains tables; some ATS parsers may read table content out of order.")
        text = "\n".join(paras + table_text)
    else:
        text = data.decode("utf-8", errors="ignore")
    if len(text.strip()) < 250: issues.append("Very little text was extracted; parsing confidence is reduced.")
    return text, issues

def _section_lines(text, aliases):
    lines = [x.strip() for x in text.splitlines() if x.strip()]
    start = None
    for i, line in enumerate(lines):
        if line.lower().strip(":") in aliases: start = i + 1; break
    if start is None: return []
    known = {a for vals in SECTION_ALIASES.values() for a in vals}
    out=[]
    for line in lines[start:]:
        if line.lower().strip(":") in known: break
        out.append(line)
    return out

def parse_resume(uploaded) -> ResumeProfile:
    text, issues = extract_text(uploaded)
    lines=[x.strip() for x in text.splitlines() if x.strip()]
    email_m=re.search(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}', text)
    phone_m=re.search(r'(?:\+?1[-.\s]?)?\(?\d{3}\)?[-.\s]\d{3}[-.\s]\d{4}', text)
    gpa_m=re.search(r'GPA\s*[:\-]?\s*([0-4](?:\.\d{1,2})?)', text, re.I)
    grad_m=re.search(r'(?:graduat(?:ion|ing)|expected)\s*[:\-]?\s*((?:Spring|Summer|Fall|Winter)?\s*20\d{2})', text, re.I)
    low=text.lower()
    skills=[]
    for skill, terms in SKILLS.items():
        if any(t in low for t in terms): skills.append(skill)
    # Formatting heuristics visible from extracted text / file structure.
    dates=re.findall(r'\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+20\d{2}|\b20\d{2}\b', text, re.I)
    if not dates: issues.append("No recognizable dates were found; experience chronology may be difficult to parse.")
    if len(lines) > 0 and (not email_m or not phone_m): issues.append("Contact information appears incomplete or was not reliably parsed.")
    confidence=max(35, 100 - 12*len(issues))
    return ResumeProfile(
        raw_text=text, name=lines[0] if lines else "", email=email_m.group(0) if email_m else "",
        phone=phone_m.group(0) if phone_m else "", education=_section_lines(text, SECTION_ALIASES["education"]),
        experience=_section_lines(text, SECTION_ALIASES["experience"]), skills=skills,
        certifications=_section_lines(text, SECTION_ALIASES["certifications"]), projects=_section_lines(text, SECTION_ALIASES["projects"]),
        leadership=_section_lines(text, SECTION_ALIASES["leadership"]), coursework=_section_lines(text, SECTION_ALIASES["coursework"]),
        gpa=float(gpa_m.group(1)) if gpa_m else None, graduation=grad_m.group(1) if grad_m else "",
        parsing_issues=issues, parsing_confidence=confidence)
